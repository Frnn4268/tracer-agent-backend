import base64
from collections import deque
from io import BytesIO

import httpx
from docx import Document
from fastapi import APIRouter, HTTPException
from app.models.schemas import Attachment, ChatRequest, ChatResponse, Message
from app.services.pipeline import PacketTracerPipeline
from app.services.openrouter import (
    OpenRouterClient,
    OpenRouterConfigurationError,
    OpenRouterEmptyResponseError,
    OpenRouterError,
)

router = APIRouter(prefix="/api/chat", tags=["chat"])
client = OpenRouterClient()
pipeline = PacketTracerPipeline(client)

SYSTEM_PROMPT = """Eres un instructor experto y arquitecto de redes especializado en Cisco Packet Tracer.

OBJETIVO PRINCIPAL:
Asistir en la resolución de laboratorios, análisis de errores, diseño de topologías y dudas técnicas de redes Cisco, proporcionando respuestas claras, precisas y pedagógicas.

REGLAS DE COMPORTAMIENTO Y TONO:
- Idioma: Responde siempre en español.
- Tono: Profesional, alentador y directo. Si el usuario comete un error, corrígelo con tacto y explica el "porqué" técnico detrás del error.
- Foco: Limítate estrictamente a temas de redes, protocolos y Cisco Packet Tracer. Si te preguntan algo fuera de este ámbito, declina amablemente.
- Veracidad (Cero Alucinaciones): NO inventes topologías, nombres de interfaces (ej. GigabitEthernet0/0), direcciones IP ni contraseñas. Si falta información crucial para dar un comando exacto, pide los datos mínimos necesarios.

METODOLOGÍA DE RESOLUCIÓN:
1. Procedimientos y Proyectos: Si el usuario pide configurar algo desde cero o un procedimiento largo, DIVÍDELO. 
   - Primero, entrega un plan breve de pasos enumerados.
   - Segundo, desarrolla ÚNICAMENTE el Paso 1.
   - Tercero, pide confirmación o validación antes de continuar con el siguiente paso.
2. Análisis de Errores (Troubleshooting): Si el usuario reporta una falla de conectividad o configuración, pídele proactivamente las salidas de comandos de diagnóstico pertinentes (ej. `ping`, `traceroute`, `show ip route`, `show running-config` de las interfaces afectadas) antes de adivinar el problema.
3. Explicación de Comandos: Acompaña cada bloque de código de configuración con una explicación de máximo una línea por comando indicando su propósito.

REGLAS DE FORMATO:
- Usa Markdown válido en toda la respuesta.
- Comandos y configuraciones: Deben ir SIEMPRE dentro de bloques de código (```).
- Estructura: Usa títulos cortos (##) para separar secciones.
- Listas: Usa listas numeradas (1, 2, 3) para secuencias de comandos o pasos, y viñetas (-) para observaciones, diagnósticos o validaciones.
- Concisión: Omite introducciones largas ("¡Hola! Claro que sí, estaré encantado de ayudarte con..."), relleno y tablas innecesarias. Ve directo a la solución.
"""

MAX_HISTORY_MESSAGES = 12
MAX_ATTACHMENTS_PER_MESSAGE = 4
MAX_ATTACHMENT_BYTES = 4 * 1024 * 1024
MAX_DOCUMENT_TEXT_CHARS = 12000


def decode_attachment_data(attachment: Attachment) -> bytes:
    try:
        _, encoded = attachment.data_url.split(",", 1)
    except ValueError as exc:
        raise ValueError(f"El adjunto {attachment.name} no tiene un data URL válido.") from exc

    try:
        raw_bytes = base64.b64decode(encoded)
    except Exception as exc:
        raise ValueError(f"No se pudo decodificar el adjunto {attachment.name}.") from exc

    if len(raw_bytes) > MAX_ATTACHMENT_BYTES:
        raise ValueError(f"El adjunto {attachment.name} supera el límite de 4 MB.")

    return raw_bytes


def extract_docx_text(raw_bytes: bytes) -> str:
    document = Document(BytesIO(raw_bytes))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n".join(parts).strip()


def build_user_content(message: Message) -> tuple[str | list[dict], bool]:
    if len(message.attachments) > MAX_ATTACHMENTS_PER_MESSAGE:
        raise ValueError("Solo se permiten hasta 4 adjuntos por mensaje.")

    text_parts: list[str] = []
    image_parts: list[dict] = []
    requires_vision = False

    if message.content.strip():
        text_parts.append(message.content.strip())

    for attachment in message.attachments:
        raw_bytes = decode_attachment_data(attachment)
        mime_type = attachment.mime_type.lower()
        file_name = attachment.name.lower()

        if mime_type.startswith("image/"):
            requires_vision = True
            image_parts.append(
                {
                    "type": "image_url",
                    "image_url": {"url": attachment.data_url},
                }
            )
            text_parts.append(f"Imagen adjunta: {attachment.name}")
            continue

        if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or file_name.endswith(".docx"):
            extracted_text = extract_docx_text(raw_bytes)
            if not extracted_text:
                raise ValueError(f"El documento {attachment.name} no contiene texto legible.")

            text_parts.append(
                f"Documento adjunto ({attachment.name}):\n{extracted_text[:MAX_DOCUMENT_TEXT_CHARS]}"
            )
            continue

        if mime_type.startswith("text/") or file_name.endswith((".txt", ".log", ".cfg", ".md")):
            decoded_text = raw_bytes.decode("utf-8", errors="ignore").strip()
            if not decoded_text:
                raise ValueError(f"El archivo {attachment.name} no contiene texto legible.")

            text_parts.append(
                f"Archivo adjunto ({attachment.name}):\n{decoded_text[:MAX_DOCUMENT_TEXT_CHARS]}"
            )
            continue

        raise ValueError(
            f"El tipo de archivo {attachment.name} no está soportado. Usa imágenes, DOCX, TXT, LOG o MD."
        )

    merged_text = "\n\n".join(part for part in text_parts if part).strip()

    if image_parts:
        prompt_text = merged_text or "Analiza las imágenes adjuntas y explica el problema probable."
        return [{"type": "text", "text": prompt_text}, *image_parts], requires_vision

    if not merged_text:
        merged_text = "El usuario adjuntó contexto sin texto adicional. Analiza el material disponible."

    return merged_text, requires_vision


def build_messages(request: ChatRequest) -> tuple[list[dict], bool]:
    trimmed_history = deque(request.messages, maxlen=MAX_HISTORY_MESSAGES)
    messages = [{"role": "system", "content": SYSTEM_PROMPT}]
    requires_vision = False

    for msg in trimmed_history:
        if msg.role == "system":
            continue

        if msg.role == "user":
            content, message_requires_vision = build_user_content(msg)
            requires_vision = requires_vision or message_requires_vision
            messages.append({"role": msg.role, "content": content})
            continue

        messages.append({"role": msg.role, "content": msg.content})

    return messages, requires_vision

@router.post("/", response_model=ChatResponse)
async def chat(request: ChatRequest):
    try:
        messages, requires_vision = build_messages(request)
        result = await pipeline.run(
            messages,
            requires_vision=requires_vision,
            web_research_mode=request.options.web_research_mode,
        )
        return ChatResponse(reply=result.reply, metadata=result.metadata)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except OpenRouterConfigurationError as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc
    except OpenRouterEmptyResponseError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    except httpx.HTTPStatusError as exc:
        detail = exc.response.text or str(exc)
        raise HTTPException(status_code=502, detail=f"OpenRouter devolvió un error HTTP: {detail}") from exc
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=502, detail=f"No se pudo contactar con OpenRouter: {exc}") from exc
    except OpenRouterError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Error interno inesperado: {exc}") from exc